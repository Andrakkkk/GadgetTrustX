import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_styled_table(doc, rows, cols, col_widths=None):
    table = doc.add_table(rows=rows, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="D3D3D3"/>
            <w:left w:val="none"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="A0A0A0"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)
    return table

def add_header_row(table, headers, bg_color="2B6CB0"):
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], bg_color)
        set_cell_margins(hdr_cells[i], 120, 120, 160, 160)
        p = hdr_cells[i].paragraphs[0]
        if len(p.runs) > 0:
            p.runs[0].font.bold = True
            p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

def add_body_rows(table, data, alternate_color="F8FAFC"):
    for row_idx, row_data in enumerate(data):
        cells = table.rows[row_idx + 1].cells
        for col_idx, val in enumerate(row_data):
            cells[col_idx].text = str(val)
            set_cell_margins(cells[col_idx], 100, 100, 140, 140)
            if row_idx % 2 == 1:
                set_cell_background(cells[col_idx], alternate_color)
        cells[0].paragraphs[0].runs[0].font.bold = True

def add_req_block(doc, req_id, title, desc, actor, pre_cond, main_flow, alt_flow, post_cond, biz_rules, trace):
    h3 = doc.add_heading(f"{req_id}: {title}", level=3)
    h3.paragraph_format.space_before = Pt(12)
    h3.paragraph_format.space_after = Pt(4)
    
    table_data = [
        ("ID Kebutuhan", req_id),
        ("Nama Kebutuhan", title),
        ("Deskripsi", desc),
        ("Aktor Utama", actor),
        ("Pre-Kondisi", pre_cond),
        ("Alur Utama (Main Flow)", main_flow),
        ("Alur Alternatif / Eksepsi", alt_flow),
        ("Post-Kondisi", post_cond),
        ("Aturan Bisnis (Business Rules)", biz_rules),
        ("Ketertelusuran (Traceability)", trace)
    ]
    
    table = create_styled_table(doc, len(table_data) + 1, 2)
    add_header_row(table, ["Komponen Spesifikasi", "Keterangan Rinci"], bg_color="1A365D")
    add_body_rows(table, table_data)
    
    # Adjust widths roughly
    for row in table.rows:
        row.cells[0].width = Inches(2.2)
        row.cells[1].width = Inches(4.3)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

def generate_srs():
    doc = docx.Document()
    
    # Page Margins 1 inch
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    
    # ==================== COVER PAGE ====================
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(24)
    p_title.paragraph_format.space_after = Pt(4)
    run_t = p_title.add_run("SPESIFIKASI KEBUTUHAN PERANGKAT LUNAK\n(Software Requirements Specification - SRS)")
    run_t.font.size = Pt(20)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(20)
    run_sub = p_sub.add_run("Sistem Marketplace Gadget Terverifikasi — GadgetTrustX\nStandar Acuan: ISO/IEC/IEEE 29148:2018 / IEEE 830")
    run_sub.font.size = Pt(13)
    run_sub.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)
    
    # Meta table
    meta_info = [
        ("Identifikasi Dokumen", "SRS-GADGETTRUSTX-MOD2-202310370311421"),
        ("Nama Aplikasi", "GadgetTrustX (Smart Device Marketplace)"),
        ("Versi Dokumen", "v1.0 (Persiapan Modul 2)"),
        ("Fokus Peran Pengujian", "Pembeli (Buyer)"),
        ("Nama Praktikan", "Leandra Chelsea Geovani Karyono"),
        ("NIM", "202310370311421"),
        ("Kelas Praktikum", "PKPL -B"),
        ("Tanggal Terbit", "09 Oktober 2026"),
        ("Status", "Final untuk Modul 2"),
        ("Dokumen Acuan Terkait", "TP-GADGETTRUSTX-MOD1-202310370311421 (Test Plan Modul 1)")
    ]
    
    t_meta = create_styled_table(doc, len(meta_info) + 1, 2)
    add_header_row(t_meta, ["Atribut Proyek", "Keterangan Spesifikasi"], bg_color="2B6CB0")
    add_body_rows(t_meta, meta_info)
    for row in t_meta.rows:
        row.cells[0].width = Inches(2.2)
        row.cells[1].width = Inches(4.3)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(16)
    
    # Revision History
    h_rev = doc.add_heading("Riwayat Revisi Dokumen", level=2)
    h_rev.paragraph_format.space_before = Pt(14)
    h_rev.paragraph_format.space_after = Pt(6)
    
    rev_data = [
        ("1.0", "09 Oktober 2026", "Leandra Chelsea Geovani Karyono", "Penyusunan awal spesifikasi kebutuhan lengkap (12 Fitur In-Scope) untuk persiapan Modul 2")
    ]
    t_rev = create_styled_table(doc, len(rev_data) + 1, 4)
    add_header_row(t_rev, ["Versi", "Tanggal", "Penyusun", "Deskripsi Perubahan"], bg_color="2B6CB0")
    add_body_rows(t_rev, rev_data)
    
    doc.add_page_break()
    
    # ==================== DAFTAR ISI ====================
    h_toc = doc.add_heading("DAFTAR ISI", level=1)
    h_toc.paragraph_format.space_after = Pt(10)
    
    toc_items = [
        ("1. Pendahuluan", "3"),
        ("   1.1 Tujuan Dokumen", "3"),
        ("   1.2 Lingkup Produk", "3"),
        ("   1.3 Definisi, Akronim, dan Singkatan", "3"),
        ("   1.4 Dokumen Referensi", "4"),
        ("   1.5 Gambaran Umum Dokumen", "4"),
        ("2. Deskripsi Keseluruhan", "5"),
        ("   2.1 Perspektif Produk", "5"),
        ("   2.2 Fungsi Produk Utama", "5"),
        ("   2.3 Karakteristik Pengguna", "6"),
        ("   2.4 Batasan-Batasan Sistem", "6"),
        ("   2.5 Asumsi dan Ketergantungan", "6"),
        ("3. Kebutuhan Antarmuka Eksternal", "7"),
        ("   3.1 Antarmuka Pengguna (User Interface)", "7"),
        ("   3.2 Antarmuka Perangkat Lunak (Software Interface)", "7"),
        ("   3.3 Antarmuka Komunikasi (Communication Interface)", "7"),
        ("4. Kebutuhan Fungsional (Functional Requirements)", "8"),
        ("   4.1 REQ-F-01: Autentikasi Pengguna (Login)", "8"),
        ("   4.2 REQ-F-02: Registrasi Akun Pembeli (Register)", "9"),
        ("   4.3 REQ-F-03: Pemulihan Kata Sandi (Forgot Password)", "10"),
        ("   4.4 REQ-F-04: Jelajah Katalog & Pencarian Produk", "11"),
        ("   4.5 REQ-F-05: Filter Multi-Parameter Produk", "12"),
        ("   4.6 REQ-F-06: Keranjang Belanja (Shopping Cart)", "13"),
        ("   4.7 REQ-F-07: Checkout dan Pembayaran Terpadu (Midtrans)", "14"),
        ("   4.8 REQ-F-08: Verifikasi Keaslian IMEI (TrustX Scanner)", "15"),
        ("   4.9 REQ-F-09: Smart Price Checker (Valuasi Harga Pasar AI)", "16"),
        ("   4.10 REQ-F-10: Smart Matching Rekomendasi Cerdas", "17"),
        ("   4.11 REQ-F-11: Ulasan dan Penilaian Produk (Review & Rating)", "18"),
        ("   4.12 REQ-F-12: Pengajuan Retur & Komplain Kerusakan Barang", "19"),
        ("5. Kebutuhan Non-Fungsional (Non-Functional Requirements)", "20"),
        ("   5.1 Kinerja dan Efisiensi (Performance)", "20"),
        ("   5.2 Keamanan dan Integritas Data (Security)", "20"),
        ("   5.3 Keandalan dan Ketersediaan (Reliability & Availability)", "20"),
        ("   5.4 Kemudahan Penggunaan (Usability)", "20"),
        ("   5.5 Portabilitas dan Kompatibilitas Browser", "20"),
        ("6. Matriks Ketertelusuran Kebutuhan (Requirements Traceability Matrix - RTM)", "21"),
        ("Lampiran A: Glosarium Istilah SQA & Standar IEEE", "22")
    ]
    
    for text, page in toc_items:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.font.size = Pt(10)
        p.add_run(f"    {page}")
        
    doc.add_page_break()
    
    # ==================== BAB 1 ====================
    doc.add_heading("1. Pendahuluan", level=1)
    
    doc.add_heading("1.1 Tujuan Dokumen", level=2)
    p = doc.add_paragraph()
    p.add_run("Dokumen Spesifikasi Kebutuhan Perangkat Lunak (Software Requirements Specification - SRS) ini disusun sebagai dasar formal spesifikasi kebutuhan fungsional dan non-fungsional untuk platform ")
    r_b = p.add_run("GadgetTrustX (Smart Device Marketplace)")
    r_b.bold = True
    p.add_run(". Dokumen ini menjadi acuan utama (Test Basis) untuk pelaksanaan praktikum Pengujian Kualitas Perangkat Lunak (PKPL) pada Modul 2 dan modul-modul berikutnya, khususnya dalam merancang kasus uji (Test Case), matriks penelusuran (Traceability Matrix), dan evaluasi kesesuaian sistem.")
    
    doc.add_heading("1.2 Lingkup Produk", level=2)
    p = doc.add_paragraph()
    p.add_run("GadgetTrustX adalah platform marketplace perangkat pintar (smart device) terintegrasi berbasis web yang dirancang untuk memfasilitasi transaksi jual-beli gadget baru dan bekas secara aman, transparan, dan terpercaya. Sistem dilengkapi fitur verifikasi keaslian IMEI melalui database TAC global, valuasi estimasi harga wajar pasar berbasis Gemini AI, rekomendasi spesifikasi cerdas (Smart Matching), serta sistem transaksi aman menggunakan escrow pembayaran Midtrans Payment Gateway.")

    doc.add_heading("1.3 Definisi, Akronim, dan Singkatan", level=2)
    def_data = [
        ("SRS", "Software Requirements Specification (Spesifikasi Kebutuhan Perangkat Lunak) berdasarkan standar ISO/IEC/IEEE 29148:2018."),
        ("IMEI", "International Mobile Equipment Identity; nomor unik 15 digit pengenal identitas perangkat seluler global."),
        ("TAC", "Type Allocation Code; 8 digit pertama nomor IMEI yang mengidentifikasi pabrikan dan model perangkat."),
        ("Algoritma Luhn", "Formula checksum modulus 10 yang digunakan untuk memvalidasi keabsahan susunan angka IMEI."),
        ("Midtrans", "Penyedia gerbang pembayaran (payment gateway) resmi untuk transaksi QRIS, transfer bank, dan e-wallet."),
        ("Supabase", "BaaS (Backend-as-a-Service) sumber terbuka berbasis PostgreSQL yang mengelola database, autentikasi, dan media storage."),
        ("RTM", "Requirements Traceability Matrix; tabel pelacak keterhubungan antara kebutuhan SRS, risiko, dan test case.")
    ]
    t_def = create_styled_table(doc, len(def_data) + 1, 2)
    add_header_row(t_def, ["Istilah / Akronim", "Definisi Operasional"], bg_color="2B6CB0")
    add_body_rows(t_def, def_data)
    for r in t_def.rows:
        r.cells[0].width = Inches(1.8)
        r.cells[1].width = Inches(4.7)

    doc.add_heading("1.4 Dokumen Referensi", level=2)
    p = doc.add_paragraph()
    p.add_run("1. ISO/IEC/IEEE 29148:2018 - Systems and software engineering — Life cycle processes — Requirements engineering.\n")
    p.add_run("2. ISO/IEC/IEEE 29119:2022 - Software and systems engineering — Software testing (Part 1, 2, 3).\n")
    p.add_run("3. Dokumen Rencana Pengujian Modul 1 (Test Plan): TP-GADGETTRUSTX-MOD1-202310370311421.")

    doc.add_heading("1.5 Gambaran Umum Dokumen", level=2)
    p = doc.add_paragraph()
    p.add_run("Dokumen ini terbagi menjadi 6 bab utama. Bab 2 memaparkan deskripsi umum sistem dan karakteristik peran Pembeli. Bab 3 menjelaskan antarmuka eksternal. Bab 4 mendefinisikan 12 kebutuhan fungsional secara rinci. Bab 5 menjabarkan kebutuhan non-fungsional (performa, keamanan, dll). Bab 6 menyajikan Matriks Ketertelusuran Kebutuhan (RTM).")

    doc.add_page_break()

    # ==================== BAB 2 ====================
    doc.add_heading("2. Deskripsi Keseluruhan", level=1)
    
    doc.add_heading("2.1 Perspektif Produk", level=2)
    p = doc.add_paragraph()
    p.add_run("GadgetTrustX merupakan aplikasi web modern mandiri (self-contained web application) yang dibangun menggunakan arsitektur Next.js 16 App Router, React 19, dan Tailwind CSS v4 di sisi frontend, berkomunikasi langsung dengan Supabase PostgreSQL dan Supabase Auth via Secure Route Handlers, serta terhubung dengan layanan pihak ketiga (Midtrans Payment Gateway dan Gemini AI Engine).")

    doc.add_heading("2.2 Fungsi Produk Utama", level=2)
    p = doc.add_paragraph()
    p.add_run("Sistem menyediakan serangkaian fitur utama yang dikhususkan untuk memenuhi kebutuhan transaksi pembeli:\n")
    p.add_run("• Pengelolaan akun dan kredensial aman (Login, Register, Forgot Password).\n")
    p.add_run("• Eksplorasi katalog gadget cerdas dengan pencarian teks bebas dan filter multi-parameter (RAM, Storage hingga 2TB, rentang harga, kondisi).\n")
    p.add_run("• Validasi keaslian perangkat (TrustX Scanner IMEI berbasis TAC & Luhn).\n")
    p.add_run("• Asistensi pemilihan gadget berbasis AI (Smart Price Checker & Smart Matching).\n")
    p.add_run("• Siklus transaksi belanja aman (Keranjang Belanja, Kalkulasi Ongkir Otomatis, Checkout Escrow Midtrans).\n")
    p.add_run("• Perlindungan pasca-pembelian (Ulasan/Rating Produk dan Formulir Pengajuan Retur Komplain Barang).")

    doc.add_heading("2.3 Karakteristik Pengguna", level=2)
    p = doc.add_paragraph()
    p.add_run("Fokus pengujian praktikum ini ditujukan secara eksklusif pada peran ")
    r_b = p.add_run("Pembeli (Buyer)")
    r_b.bold = True
    p.add_run(". Pembeli adalah konsumen akhir yang memiliki keterampilan digital tingkat menengah, berinteraksi melalui browser desktop atau mobile untuk mencari, mengecek keaslian, membeli gadget, dan mengelola tiket komplain pesanan.")

    doc.add_heading("2.4 Batasan-Batasan Sistem", level=2)
    p = doc.add_paragraph()
    p.add_run("1. Pengujian dibatasi khusus pada antarmuka dan interaksi peran Pembeli; fitur manajemen stok toko oleh Penjual (Seller) dan dasbor operasional Administrator berada di luar cakupan.\n")
    p.add_run("2. Transaksi finansial pada pembayaran checkout diuji menggunakan lingkungan Midtrans Sandbox untuk menghindari mutasi uang riil.\n")
    p.add_run("3. Estimasi valuasi harga pada Smart Price Checker bersumber dari model AI Gemini dan berfungsi sebagai referensi pasar dinamis.")

    doc.add_heading("2.5 Asumsi dan Ketergantungan", level=2)
    p = doc.add_paragraph()
    p.add_run("• Pengguna memiliki koneksi internet stabil saat mengakses web.\n")
    p.add_run("• Basis data Supabase menyediakan dataset produk aktif sebelum pengujian berlangsung.\n")
    p.add_run("• Layanan pihak ketiga (Midtrans Sandbox dan Supabase Auth) beroperasi dalam kondisi normal.")

    doc.add_page_break()

    # ==================== BAB 3 ====================
    doc.add_heading("3. Kebutuhan Antarmuka Eksternal", level=1)
    
    doc.add_heading("3.1 Antarmuka Pengguna (User Interface)", level=2)
    p = doc.add_paragraph()
    p.add_run("Antarmuka pengguna dirancang responsif (desktop, tablet, mobile), mengusung tema modern dengan dukungan peralihan Dark Mode dan Light Mode. Setiap komponen input dilengkapi validasi interaktif, batasan karakter, serta pesan kesalahan yang jelas.")

    doc.add_heading("3.2 Antarmuka Perangkat Lunak (Software Interface)", level=2)
    p = doc.add_paragraph()
    p.add_run("• Database Server: PostgreSQL via Supabase SDK (Client-Server REST/RPC).\n")
    p.add_run("• Payment Gateway: Midtrans Snap API (Sandbox Environment) untuk pembuatan token transaksi pembayaran.\n")
    p.add_run("• AI Valuation Engine: Google Gemini API untuk analisis harga pasar gadget.")

    doc.add_heading("3.3 Antarmuka Komunikasi (Communication Interface)", level=2)
    p = doc.add_paragraph()
    p.add_run("Seluruh pertukaran data antara browser klien dan server menggunakan protokol aman HTTPS (TLS 1.3) dan REST API dengan format payload JSON.")

    doc.add_page_break()

    # ==================== BAB 4 ====================
    doc.add_heading("4. Kebutuhan Fungsional (Functional Requirements)", level=1)
    p_intro = doc.add_paragraph()
    p_intro.add_run("Bagian ini merinci 12 Kebutuhan Fungsional (REQ-F-01 s/d REQ-F-12) yang mencakup seluruh fitur In-Scope pada dokumen Test Plan Modul 1.")

    # 4.1 REQ-F-01
    add_req_block(
        doc,
        "REQ-F-01",
        "Autentikasi Pengguna (Login)",
        "Sistem memungkinkan pembeli terdaftar untuk masuk ke dalam aplikasi menggunakan alamat email dan kata sandi yang valid.",
        "Pembeli (Buyer)",
        "Akun pembeli telah terdaftar sebelumnya di database Supabase Auth.",
        "1. Pembeli membuka halaman /login.\n2. Pembeli memasukkan email dan password.\n3. Pembeli menekan tombol 'Masuk'.\n4. Sistem memverifikasi kredensial via Supabase Auth.\n5. Kredensial valid, token sesi disimpan, pembeli dialihkan ke halaman utama/dashboard.",
        "• Jika email belum terdaftar atau password salah, sistem menampilkan notifikasi kesalahan: 'Kredensial tidak valid'.\n• Jika format email salah, form menampilkan validasi instan.",
        "Sesi pengguna aktif, status login tersimpan di state global dan localStorage.",
        "Password minimal 6 karakter. Proteksi brute-force aktif via Supabase rate-limiting.",
        "Mendukung Fitur No. 1 pada Test Plan. Mitigasi proteksi autentikasi dasar."
    )

    # 4.2 REQ-F-02
    add_req_block(
        doc,
        "REQ-F-02",
        "Registrasi Akun Pembeli (Register)",
        "Sistem memungkinkan pengguna baru untuk membuat akun pembeli dengan mengisikan nama lengkap, alamat email valid, nomor telepon, dan kata sandi baru.",
        "Pengguna Baru / Pembeli",
        "Pengguna belum memiliki akun dengan email yang sama di sistem.",
        "1. Pengguna membuka halaman /register.\n2. Pengguna mengisi nama, email, no HP, dan password.\n3. Pengguna menekan tombol 'Daftar Sekarang'.\n4. Sistem memvalidasi kelengkapan data dan keunikan email.\n5. Sistem membuat entri akun baru di Supabase Auth dan profil pengguna.\n6. Pengguna otomatis diarahkan ke halaman login atau beranda.",
        "• Jika email sudah pernah terdaftar, sistem menampilkan pesan error: 'Email sudah digunakan'.\n• Jika konfirmasi password tidak cocok, form memblokir pengiriman.",
        "Akun pembeli baru berhasil tercatat di database dan dapat digunakan untuk login.",
        "Role pengguna otomatis ditetapkan sebagai 'buyer'. Batas input teks nama maksimal 100 karakter.",
        "Mendukung Fitur No. 2 pada Test Plan."
    )

    # 4.3 REQ-F-03
    add_req_block(
        doc,
        "REQ-F-03",
        "Pemulihan Kata Sandi (Forgot Password)",
        "Sistem menyediakan alur permintaan tautan reset password yang dikirimkan ke alamat email terdaftar.",
        "Pembeli (Buyer)",
        "Akun pembeli telah terdaftar di sistem.",
        "1. Pembeli mengakses tautan 'Lupa Password' di /login.\n2. Pembeli memasukkan alamat email terdaftar pada form /reset-password.\n3. Pembeli menekan tombol 'Kirim Tautan Pemulihan'.\n4. Sistem mengirimkan email berisi tautan token reset password via Supabase Auth.\n5. Sistem menampilkan notifikasi sukses pengiriman.",
        "• Jika format email tidak valid, sistem menampilkan peringatan format.\n• Jika email tidak ditemukan, sistem menampilkan pesan netral demi alasan keamanan privasi.",
        "Tautan pemulihan terkirim ke inbox email pengguna dengan masa berlaku 1 jam.",
        "Tautan hanya dapat digunakan satu kali (one-time token).",
        "Mendukung Fitur No. 3 pada Test Plan."
    )

    # 4.4 REQ-F-04
    add_req_block(
        doc,
        "REQ-F-04",
        "Jelajah Katalog & Pencarian Produk",
        "Sistem menyediakan katalog gadget aktif dengan fungsi pencarian kata kunci berdasarkan nama produk, model, atau merek.",
        "Pembeli (Buyer)",
        "Tersedia data listing produk gadget aktif pada database.",
        "1. Pembeli membuka halaman /marketplace.\n2. Pembeli mengetik kata kunci pada search bar (misal 'iPhone' atau 'Samsung').\n3. Sistem memfilter daftar produk secara real-time berdasarkan kecocokan nama, brand, atau deskripsi.\n4. Sistem menampilkan kartu produk yang sesuai beserta harga, spesifikasi, dan rating.",
        "• Jika kata kunci tidak cocok dengan produk apa pun, sistem menampilkan tampilan 'Tidak ada produk ditemukan' disertai tombol reset.",
        "Daftar produk yang ditampilkan sesuai dengan kata kunci pencarian.",
        "Panjang input kata kunci dibatasi maksimal 100 karakter untuk mencegah injection.",
        "Mendukung Fitur No. 4 pada Test Plan."
    )

    # 4.5 REQ-F-05
    add_req_block(
        doc,
        "REQ-F-05",
        "Filter Multi-Parameter Produk",
        "Sistem menyediakan fitur penyaringan katalog berdasarkan banyak kriteria sekaligus: Brand, RAM, Internal Storage (hingga opsi 2TB), kondisi, dan rentang harga.",
        "Pembeli (Buyer)",
        "Pembeli berada di halaman katalog /marketplace.",
        "1. Pembeli membuka panel filter (accordion collapsible).\n2. Pembeli memilih opsi Brand, RAM (misal 8GB), Storage (misal 256GB / 2TB), dan mengisi batas harga min/max.\n3. Sistem memperbarui query filter dan menyajikan produk yang memenuhi semua kriteria.",
        "• Pembeli dapat mereset salah satu filter dengan mengklik badge tag aktif (✕), atau mereset seluruh filter sekaligus melalui tombol 'Reset Semua'.",
        "Katalog produk menampilkan unit gadget yang memenuhi kriteria gabungan.",
        "Batas input harga minimum dan maksimum hanya menerima angka positif hingga Rp 100.000.000.",
        "Mendukung Fitur No. 5 pada Test Plan. Menguji Risiko R-04 (kombinasi filter)."
    )

    # 4.6 REQ-F-06
    add_req_block(
        doc,
        "REQ-F-06",
        "Keranjang Belanja (Shopping Cart)",
        "Sistem memungkinkan pembeli menyimpan unit gadget ke keranjang, memperbarui jumlah kuantitas sesuai stok barang, dan menghapus item dari keranjang.",
        "Pembeli (Buyer)",
        "Pembeli sudah login dan memilih produk dengan status stok tersedia.",
        "1. Pembeli menekan tombol 'Tambah ke Keranjang' pada kartu produk / halaman detail.\n2. Animasi fly-to-cart berjalan dan badge keranjang bertambah.\n3. Pembeli membuka halaman /cart untuk mengubah kuantitas atau menghapus item.\n4. Sistem menghitung subtotal belanja secara otomatis.",
        "• Jika pembeli menambah kuantitas melebihi stok yang ada, sistem membatasi penambahan dan menampilkan notifikasi batas stok.\n• Akun peran Penjual diblokir untuk menambah produk miliknya sendiri ke keranjang.",
        "Data keranjang belanja tersimpan di database pengguna dan siap diteruskan ke checkout.",
        "Kuantitas minimal 1 unit, kuantitas maksimal dibatasi oleh sisa stok penjual.",
        "Mendukung Fitur No. 6 pada Test Plan."
    )

    # 4.7 REQ-F-07
    add_req_block(
        doc,
        "REQ-F-07",
        "Checkout dan Pembayaran Terpadu (Midtrans Payment)",
        "Sistem memfasilitasi transaksi checkout, pemilihan alamat tujuan, estimasi ongkir ekspedisi, pembuatan invoice token Midtrans, dan verifikasi status pembayaran lunas.",
        "Pembeli (Buyer)",
        "Keranjang belanja memiliki minimal 1 produk dan stok unit masih tersedia.",
        "1. Pembeli menekan tombol 'Checkout' dari halaman /cart.\n2. Pembeli memilih alamat kirim dan opsi kurir ekspedisi di /checkout.\n3. Sistem menghitung total tagihan = subtotal barang + ongkos kirim.\n4. Pembeli menekan 'Bayar Sekarang' $\\rightarrow$ popup Midtrans Snap muncul.\n5. Pembeli memilih metode (QRIS/Transfer VA) pada simulator sandbox dan menyelesaikan pembayaran.\n6. Webhook Midtrans memperbarui status pesanan menjadi 'Paid' dan invoice sukses dibuat.",
        "• Jika transaksi dibatalkan oleh pengguna, token transaksi kedaluwarsa dan status pesanan menjadi 'Cancelled'.\n• Jika stok habis saat proses checkout berlangsung, sistem membatalkan transaksi secara aman.",
        "Pesanan tercatat di database dengan status lunas, stok produk berkurang otomatis.",
        "Kalkulasi total harga dihitung ulang dan diverifikasi pada server-side sebelum menerbitkan snap token.",
        "Mendukung Fitur No. 7 pada Test Plan. Mitigasi Risiko Kritis R-07 (Skor 15)."
    )

    # 4.8 REQ-F-08
    add_req_block(
        doc,
        "REQ-F-08",
        "Verifikasi Keaslian IMEI (TrustX Scanner)",
        "Sistem menyediakan modul pemindai IMEI untuk memvalidasi nomor seri 15 digit menggunakan algoritma Luhn dan mencocokkannya dengan database TAC resmi.",
        "Pembeli (Buyer)",
        "Pembeli mengakses halaman /scanner.",
        "1. Pembeli memasukkan 15 digit nomor IMEI perangkat.\n2. Pembeli menekan tombol 'Periksa Keaslian IMEI'.\n3. Sistem memvalidasi checksum menggunakan algoritma Luhn.\n4. Sistem mencocokkan 8 digit pertama (TAC) dengan database referensi global.\n5. Sistem menampilkan hasil audit: Nama Model, Merek Resmi, Status Legalitas, dan Rekomendasi Keamanan.",
        "• Jika input kurang atau lebih dari 15 digit, sistem menampilkan pesan: 'Nomor IMEI harus tepat 15 digit angka'.\n• Jika input mengandung huruf/simbol, sistem menolak input secara otomatis.",
        "Hasil verifikasi ditampilkan secara transparan di layar pembeli.",
        "Input hanya menerima format numerik persis 15 digit. Checksum digit ke-15 harus valid modulo 10.",
        "Mendukung Fitur No. 8 pada Test Plan. Mitigasi Risiko R-01 (Skor 12)."
    )

    # 4.9 REQ-F-09
    add_req_block(
        doc,
        "REQ-F-09",
        "Smart Price Checker (Valuasi Harga Pasar AI)",
        "Sistem menyediakan kalkulator valuasi harga berbasis AI Gemini untuk memperkirakan nilai wajar pasar gadget berdasarkan spesifikasi, kondisi, dan kelengkapan.",
        "Pembeli (Buyer)",
        "Pembeli mengakses halaman /price-checker.",
        "1. Pembeli memilih merek, tipe gadget, kapasitas RAM & Storage, serta kondisi fisik (Mulus, Normal, Minus).\n2. Pembeli menekan tombol 'Cek Estimasi Harga AI'.\n3. Sistem memanggil Gemini AI API untuk menganalisis parameter.\n4. Sistem menyajikan hasil valuasi: Rentang Harga Pasar Wajar (Min - Max), Rekomendasi Harga Beli, dan Analisis Depresiasi.",
        "• Jika layanan AI mengalami latency/timeout, sistem beralih ke tabel referensi harga cadangan lokal (fallback database).",
        "Rekomendasi estimasi harga tampil secara lengkap di antarmuka pengguna.",
        "Kapasitas memori dan kondisi fisik divalidasi dengan batas nilai wajar sebelum diproses.",
        "Mendukung Fitur No. 9 pada Test Plan. Mitigasi Risiko R-05 (Skor 8)."
    )

    # 4.10 REQ-F-10
    add_req_block(
        doc,
        "REQ-F-10",
        "Smart Matching (Rekomendasi Cerdas Berbasis Kebutuhan)",
        "Sistem menyediakan wizard bertahap bagi pembeli untuk mendapatkan rekomendasi gadget terbaik berdasarkan pagu anggaran (budget) dan prioritas penggunaan.",
        "Pembeli (Buyer)",
        "Pembeli mengakses halaman /smart-matching.",
        "1. Langkah 1: Pembeli menentukan rentang budget melalui slider interaktif.\n2. Langkah 2: Pembeli memilih prioritas utama (Gaming Performa, Fotografi Kamera, atau Baterai Tahan Lama).\n3. Langkah 3: Pembeli memilih preferensi ekosistem/merek.\n4. Sistem memproses algoritma pencocokan skor spesifikasi.\n5. Sistem menyajikan 3 rekomendasi gadget paling ideal lengkap dengan skor kecocokan.",
        "• Pembeli dapat mengubah slider atau melompat kembali ke langkah sebelumnya melalui stepper navigasi.",
        "Daftar gadget rekomendasi ditampilkan dengan tautan langsung ke detail produk.",
        "Slider budget dibatasi antara Rp 1.000.000 hingga Rp 30.000.000.",
        "Mendukung Fitur No. 10 pada Test Plan."
    )

    # 4.11 REQ-F-11
    add_req_block(
        doc,
        "REQ-F-11",
        "Ulasan dan Penilaian Produk (Review & Rating)",
        "Sistem memungkinkan pembeli yang telah menyelesaikan transaksi untuk memberikan penilaian bintang (1-5), ulasan pengalaman, dan mengunggah foto produk.",
        "Pembeli (Buyer)",
        "Pembeli memiliki pesanan dengan status 'Completed' / 'Paid' pada /buyer-profile.",
        "1. Pembeli membuka tab Riwayat Pesanan di profil pembeli.\n2. Pembeli menekan tombol 'Beri Ulasan' pada produk terkait.\n3. Modal ulasan terbuka: pembeli memilih rating 1–5 bintang, menulis teks ulasan, dan mengunggah foto produk.\n4. Pembeli menekan tombol 'Kirim Ulasan'.\n5. Sistem menyimpan ulasan dan memperbarui rating rata-rata produk.",
        "• Jika rating bintang belum dipilih (0 bintang), tombol kirim terkunci (disabled).\n• Jika produk sudah pernah diulas, tombol ulasan otomatis berganti menjadi 'Sudah Diulas'.",
        "Ulasan tercatat di database dan tampil pada halaman detail produk publik.",
        "Teks ulasan maksimal 500 karakter. Format foto terbatas pada JPG, PNG, WEBP dengan batas ukuran 3MB.",
        "Mendukung Fitur No. 11 pada Test Plan. Mitigasi Risiko R-03 (Skor 12)."
    )

    # 4.12 REQ-F-12
    add_req_block(
        doc,
        "REQ-F-12",
        "Pengajuan Retur & Komplain Kerusakan Barang",
        "Sistem menyediakan formulir klaim retur/komplain bagi pembeli untuk melaporkan produk cacat, menyertakan alasan kerusakan, dan melampirkan foto bukti fisik.",
        "Pembeli (Buyer)",
        "Pembeli menerima barang yang bermasalah dan masa garansi komplain masih aktif.",
        "1. Pembeli membuka riwayat transaksi di /buyer-profile.\n2. Pembeli menekan tombol 'Ajukan Retur / Komplain'.\n3. Pembeli memilih kategori kendala (Cacat Fisik, Fungsi Rusak, IMEI Tidak Cocok, Barang Tidak Sesuai).\n4. Pembeli mengisi deskripsi detail kerusakan dan mengunggah foto bukti fisik cacat.\n5. Pembeli menekan 'Kirim Pengajuan'.\n6. Sistem menerbitkan nomor tiket komplain dan mengubah status pesanan menjadi 'Return Pending'.",
        "• Jika berkas foto melebihi 5MB atau berformat selain gambar, sistem menolak unggahan dan menampilkan peringatan error yang jelas.\n• Pengajuan retur dibatasi hanya pada form input pembeli; alur approval dan transfer refund berada di sisi Seller/Admin (Out-of-Scope).",
        "Tiket retur tersimpan di database dan pembeli dapat memantau status komplainnya.",
        "Deskripsi kerusakan minimal 20 karakter dan maksimal 500 karakter. Wajib melampirkan minimal 1 foto bukti fisik.",
        "Mendukung Fitur No. 12 pada Test Plan. Mitigasi Risiko R-02 (Skor 12)."
    )

    doc.add_page_break()

    # ==================== BAB 5 ====================
    doc.add_heading("5. Kebutuhan Non-Fungsional (Non-Functional Requirements)", level=1)
    
    nf_data = [
        ("REQ-NF-01", "Kinerja (Performance)", "Waktu muat halaman (Page Load Time) di bawah 2 detik pada koneksi internet standar (10 Mbps). Waktu respons API Route Handlers maksimal 800 ms."),
        ("REQ-NF-02", "Keamanan (Security)", "Kredensial login dilindungi hashing enkripsi bcrypt via Supabase Auth. Form publik dilindungi Cloudflare Turnstile CAPTCHA. Transaksi finansial dilindungi escrow token Midtrans Snap."),
        ("REQ-NF-03", "Keandalan (Reliability)", "Sistem mampu menangani kegagalan API eksternal (fallback offline database saat Gemini AI timeout). Ketersediaan sistem (Availability) minimal 99% pada platform deployment."),
        ("REQ-NF-04", "Kemudahan Penggunaan (Usability)", "Antarmuka mendukung Dark Mode dan Light Mode yang nyaman di mata. Desain antarmuka mematuhi prinsip kontras warna WCAG AA, navigasi ramah mobile (responsive drawer), dan input counter informatif."),
        ("REQ-NF-05", "Kompatibilitas (Compatibility)", "Aplikasi dapat diakses dan berfungsi optimal pada peramban web modern: Google Chrome (Desktop & Android), Mozilla Firefox, Microsoft Edge, dan Apple Safari."),
        ("REQ-NF-06", "Integritas Data (Data Integrity)", "Pencatatan transaksi checkout bersifat idempotent (mencegah double-order saat tombol ditekan berulang kali) dan menerapkan integritas referensial foreign-key di PostgreSQL.")
    ]
    
    t_nf = create_styled_table(doc, len(nf_data) + 1, 3)
    add_header_row(t_nf, ["ID Kebutuhan", "Kategori Kebutuhan", "Spesifikasi Tolok Ukur"], bg_color="2B6CB0")
    add_body_rows(t_nf, nf_data)
    for r in t_nf.rows:
        r.cells[0].width = Inches(1.5)
        r.cells[1].width = Inches(2.0)
        r.cells[2].width = Inches(3.2)

    doc.add_page_break()

    # ==================== BAB 6 ====================
    doc.add_heading("6. Matriks Ketertelusuran Kebutuhan (Requirements Traceability Matrix - RTM)", level=1)
    p = doc.add_paragraph()
    p.add_run("Matriks berikut menunjukkan hubungan ketertelusuran dua arah (traceability) antara Kebutuhan Fungsional SRS, Fitur Test Plan Modul 1, ID Risiko Terkait, dan Rencana Teknik Pengujian untuk Modul berikutnya:")

    rtm_data = [
        ("REQ-F-01", "Autentikasi Login", "Fitur 1 (Auth)", "R-06 (Waktu/Proyek)", "Black-Box (EP & BVA)", "Siap Diuji"),
        ("REQ-F-02", "Registrasi Akun", "Fitur 2 (Auth)", "R-06 (Waktu/Proyek)", "Black-Box (EP & BVA)", "Siap Diuji"),
        ("REQ-F-03", "Forgot Password", "Fitur 3 (Auth)", "R-06 (Waktu/Proyek)", "State Transition Test", "Siap Diuji"),
        ("REQ-F-04", "Pencarian Katalog", "Fitur 4 (Marketplace)", "R-04 (Pencarian/Filter)", "Equivalence Partitioning", "Siap Diuji"),
        ("REQ-F-05", "Filter Multi-Parameter", "Fitur 5 (Filter)", "R-04 (Filter Query)", "Decision Table Testing", "Siap Diuji"),
        ("REQ-F-06", "Keranjang Belanja", "Fitur 6 (Cart)", "R-07 (Subtotal Hitung)", "BVA & Boundary Test", "Siap Diuji"),
        ("REQ-F-07", "Checkout & Midtrans", "Fitur 7 (Payment)", "R-07 (Skor 15 - Kritis)", "End-to-End Workflow Test", "Prioritas Utama"),
        ("REQ-F-08", "TrustX Scanner IMEI", "Fitur 8 (Scanner)", "R-01 (Skor 12 - Format)", "BVA & Luhn Checksum Test", "Prioritas Utama"),
        ("REQ-F-09", "Smart Price Checker", "Fitur 9 (AI Valuasi)", "R-05 (Skor 8 - Estimasi)", "Combinatorial Testing", "Siap Diuji"),
        ("REQ-F-10", "Smart Matching", "Fitur 10 (Matching)", "R-05 (Rekomendasi)", "Equivalence Partitioning", "Siap Diuji"),
        ("REQ-F-11", "Ulasan Produk", "Fitur 11 (Review)", "R-03 (Skor 12 - Ganda)", "Boundary & State Test", "Prioritas Utama"),
        ("REQ-F-12", "Pengajuan Retur", "Fitur 12 (Retur)", "R-02 (Skor 12 - Foto)", "Negative & File Upload Test", "Prioritas Utama")
    ]

    t_rtm = create_styled_table(doc, len(rtm_data) + 1, 6)
    add_header_row(t_rtm, ["ID SRS", "Nama Fitur", "Ref Test Plan", "Risiko Terkait", "Teknik Rencana Uji", "Status Kesiapan"], bg_color="1A365D")
    add_body_rows(t_rtm, rtm_data)
    for r in t_rtm.rows:
        r.cells[0].width = Inches(1.1)
        r.cells[1].width = Inches(1.4)
        r.cells[2].width = Inches(1.1)
        r.cells[3].width = Inches(1.2)
        r.cells[4].width = Inches(1.2)
        r.cells[5].width = Inches(1.0)

    doc.add_page_break()

    # ==================== LAMPIRAN A ====================
    doc.add_heading("Lampiran A: Glosarium Istilah SQA & Standar IEEE", level=1)
    
    glo_data = [
        ("Requirement (Kebutuhan)", "Kondisi atau kemampuan yang harus dipenuhi atau dimiliki oleh sistem untuk memuaskan kontrak, standar, atau kebutuhan pemangku kepentingan."),
        ("Functional Requirement", "Kebutuhan yang menspesifikasikan fungsi atau aksi yang harus mampu dieksekusi oleh sistem (apa yang dilakukan sistem)."),
        ("Non-Functional Requirement", "Kebutuhan yang menspesifikasikan kriteria kualitas, batasan, atau atribut performa sistem (seberapa baik sistem melakukannya)."),
        ("Test Basis", "Seluruh dokumen atau sumber pengetahuan yang menjadi dasar dalam menurunkan kondisi uji dan kasus uji."),
        ("Traceability Matrix", "Tabel yang menghubungkan kebutuhan bisnis ke kebutuhan perangkat lunak dan ke kasus uji untuk memastikan tidak ada kebutuhan yang terlewat."),
        ("Equivalence Partitioning (EP)", "Teknik pengujian black-box yang membagi domain input ke dalam kelas-kelas data ekuivalen yang diasumsikan diproses secara sama."),
        ("Boundary Value Analysis (BVA)", "Teknik pengujian black-box yang menguji nilai pada batas ekstrem (tepat di batas, tepat di bawah, dan tepat di atas batas partisi).")
    ]
    t_glo = create_styled_table(doc, len(glo_data) + 1, 2)
    add_header_row(t_glo, ["Istilah Kunci", "Penjelasan Standar ISO/IEC/IEEE"], bg_color="2B6CB0")
    add_body_rows(t_glo, glo_data)
    for r in t_glo.rows:
        r.cells[0].width = Inches(2.2)
        r.cells[1].width = Inches(4.5)

    doc.save("Dokumen_SRS_GadgetTrustX_Modul2.docx")
    print("Dokumen_SRS_GadgetTrustX_Modul2.docx successfully created!")

if __name__ == "__main__":
    generate_srs()
