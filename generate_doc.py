import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_document():
    doc = docx.Document()

    # Page Margins: Normal 1 inch (2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base styling
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    def set_cell_background(cell, hex_color):
        shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
        cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

    def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for margin, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{margin}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    def add_hyperlink(paragraph, url, text, color="0563C1", underline=True):
        part = paragraph.part
        r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        hyperlink = parse_xml(
            f'<w:hyperlink {nsdecls("w")} {nsdecls("r")} r:id="{r_id}">'
            f'<w:r>'
            f'<w:rPr>'
            f'<w:color w:val="{color}"/>'
            f'{"<w:u w:val=\'single\'/>" if underline else ""}'
            f'</w:rPr>'
            f'<w:t>{text}</w:t>'
            f'</w:r>'
            f'</w:hyperlink>'
        )
        paragraph._p.append(hyperlink)

    def set_table_borders(table, color="B0C4DE", sz="6", val="single"):
        tblPr = table._tbl.tblPr
        tblBorders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(tblBorders)

    def add_custom_heading(text, level, space_before=14, space_after=6):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Calibri'
        if level == 1:
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D) # Navy
        elif level == 2:
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        elif level == 3:
            run.font.size = Pt(11.5)
            run.font.color.rgb = RGBColor(0x2C, 0x52, 0x82) # Slate Blue
        return p

    # --- TITLE SECTION ---
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    r1 = p_title.add_run("Spesifikasi Kebutuhan Perangkat Lunak (SKPL / Software Requirements Specification - SRS)\n")
    r1.font.name = 'Calibri'
    r1.font.size = Pt(13)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    r2 = p_title.add_run("User Story & Requirement Document")
    r2.font.name = 'Calibri'
    r2.font.size = Pt(22)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(2)
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run("Platform GadgetTrustX: Marketplace Gadget Terverifikasi, AI Trade-In, Scanner IMEI & Escrow Terpadu")
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(12)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)

    # --- METADATA TABLE ---
    meta_table = doc.add_table(rows=8, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(meta_table, color="B0C4DE", sz="6")
    
    meta_rows = [
        ("Atribut Dokumen", "Keterangan Spesifikasi", "header"),
        ("Penyusun", "[Nama Lengkap Anda]  (NIM: [NIM Anda])", "text"),
        ("Target Pembaca", "Business Analyst (BA), Software Architect, QA Engineer & Developer", "text"),
        ("Versi Dokumen", "1.0 — Rilis Awal: Matriks Skenario Tanpa Kolom 'Then'", "text"),
        ("Status Sistem", "Platform Marketplace Gadget — Ter-deploy di Netlify (Production)", "text"),
        ("Repositori Kode", "https://github.com/Andrakkkk/GadgetTrustX", "link_github"),
        ("URL Deploy", "https://gadgetrustx.netlify.app/", "link_deploy"),
        ("Tanggal Pembaruan", "29 September 2026", "text")
    ]

    col_widths = [Inches(2.0), Inches(4.5)]
    for r_idx, (k, v, row_type) in enumerate(meta_rows):
        row = meta_table.rows[r_idx]
        cell_k = row.cells[0]
        cell_v = row.cells[1]
        
        cell_k.width = col_widths[0]
        cell_v.width = col_widths[1]
        
        set_cell_margins(cell_k, top=100, bottom=100, left=140, right=140)
        set_cell_margins(cell_v, top=100, bottom=100, left=140, right=140)
        
        p_k = cell_k.paragraphs[0]
        p_k.paragraph_format.space_before = Pt(0)
        p_k.paragraph_format.space_after = Pt(0)
        
        p_v = cell_v.paragraphs[0]
        p_v.paragraph_format.space_before = Pt(0)
        p_v.paragraph_format.space_after = Pt(0)
        
        if row_type == "header":
            set_cell_background(cell_k, "1B365D")
            set_cell_background(cell_v, "1B365D")
            
            run_k = p_k.add_run(k)
            run_k.font.name = 'Calibri'
            run_k.font.size = Pt(10)
            run_k.bold = True
            run_k.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            
            run_v = p_v.add_run(v)
            run_v.font.name = 'Calibri'
            run_v.font.size = Pt(10)
            run_v.bold = True
            run_v.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        else:
            set_cell_background(cell_k, "F0F4F8")
            set_cell_background(cell_v, "FFFFFF")
            
            run_k = p_k.add_run(k)
            run_k.font.name = 'Calibri'
            run_k.font.size = Pt(10)
            run_k.bold = True
            run_k.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
            
            if row_type == "link_github":
                add_hyperlink(p_v, "https://github.com/Andrakkkk/GadgetTrustX", "https://github.com/Andrakkkk/GadgetTrustX", color="0563C1", underline=True)
            elif row_type == "link_deploy":
                add_hyperlink(p_v, "https://gadgetrustx.netlify.app/", "https://gadgetrustx.netlify.app/", color="0563C1", underline=True)
            else:
                run_v = p_v.add_run(v)
                run_v.font.name = 'Calibri'
                run_v.font.size = Pt(10)
                run_v.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(6)
    p_spacer.paragraph_format.space_after = Pt(6)

    # --- BAB 1 ---
    add_custom_heading("1. Pendahuluan & Gambaran Umum Sistem", level=1)
    
    p_intro = doc.add_paragraph(
        "GadgetTrustX merupakan platform marketplace smart device (smartphone, tablet, laptop, dan wearable) "
        "terpadu yang dirancang khusus untuk memecahkan problematika tingginya risiko penipuan pada transaksi jual-beli gadget bekas "
        "di Indonesia. Permasalahan utama seperti barang fiktif, unit curian/blokir IMEI, manipulasi kondisi kesehatan baterai (Battery Health), "
        "serta ketidakpastian transfer langsung diselesaikan melalui arsitektur komprehensif yang mengintegrasikan: "
        "Rekening Bersama Terproteksi (TrustX Escrow Payment), Mesin Valuasi & Tukar Tambah AI (Gemini AI Trade-In Engine), "
        "Pemindai Legalitas IMEI TAC Global (TrustX Scanner), Rekomendasi Pintar (Smart Match), serta Negosiasi Langsung (Live Chat & Custom Offer)."
    )
    
    add_custom_heading("1.1 Arsitektur Marketplace Terpercaya & 5 Modul Utama Navigasi", level=2)
    p_arch = doc.add_paragraph(
        "Platform GadgetTrustX dapat diakses secara daring melalui peramban web pada domain resmi https://gadgetrustx.netlify.app. "
        "Antarmuka platform disusun dalam tema futuristik (Dark Glassmorphism) berbasis Tailwind CSS dan Next.js App Router, "
        "yang memfasilitasi 5 pilar fungsionalitas utama pada bilah navigasi (Navbar):"
    )

    modules = [
        ("Marketplace (/marketplace)", "Katalog komprehensif produk gadget baru dan bekas dengan sistem filter multi-parameter (Brand, RAM, Storage, Kategori, Rentang Harga) serta pengurutan dinamis (Termurah, Termahal, Terbaru)."),
        ("AI Valuation (/price-checker)", "Mesin valuasi harga pasar real-time bertenaga Google Gemini AI yang menghitung estimasi nilai wajar, rentang harga pasar (price range), confidence score, dan analisis tingkat permintaan pasar."),
        ("Scanner (/verification)", "Alat pemindai dan validasi 15-digit nomor IMEI dengan algoritma Luhn Check dan integrasi database TAC global (80+ model) untuk memverifikasi keaslian, spesifikasi pabrikan, status blacklist/hilang, dan garansi regional."),
        ("Smart Match (/smart-matching)", "Wizard interaktif 4-tahap yang mencocokkan kebutuhan konsumen (budget, tujuan utama seperti gaming/fotografi/produktivitas, dan preferensi merek) dengan katalog aktif menggunakan penilaian AI match score."),
        ("HP Bekas & Trade-In Hub (/trade-in)", "Pusat tukar tambah gadget digital yang menyediakan form pengunggahan 6 sudut foto kondisi fisik (Depan, Belakang, Layar, Samping, Box/Charger, Cacat/Dent) dengan potongan harga instan yang diaudit oleh AI Gemini.")
    ]

    for m_name, m_desc in modules:
        p_m = doc.add_paragraph(style='List Bullet')
        p_m.paragraph_format.space_before = Pt(2)
        p_m.paragraph_format.space_after = Pt(2)
        r_mn = p_m.add_run(f"{m_name}: ")
        r_mn.bold = True
        r_mn.font.name = 'Calibri'
        r_mn.font.size = Pt(10)
        r_mn.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        r_md = p_m.add_run(m_desc)
        r_md.font.name = 'Calibri'
        r_md.font.size = Pt(10)

    add_custom_heading("1.2 Matriks Peran & Aktor Sistem", level=2)
    
    actor_table = doc.add_table(rows=4, cols=3)
    actor_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(actor_table, color="B0C4DE", sz="6")
    
    actor_headers = ["Aktor / Persona", "Peran Utama", "Fokus Operasional & Kebutuhan dalam Aplikasi"]
    actor_rows = [
        ("Superadmin / TrustX Inspector", "Pengawas & Verifikator Platform Pusat", 
         "Mengakses dasbor admin (/admin), memantau statistik transaksi & escrow, menyetujui/menolak pengajuan retur pesanan (Return Refund Approval), memvalidasi perangkat untuk pemberian badge 'Verified by TrustX', serta mengelola akun pengguna dan badge reputasi toko."),
        ("Penjual (Seller / Toko Mitra)", "Penyedia Gadget & Pemroses Pesanan", 
         "Mengakses dasbor toko (/seller-profile), menambah produk menggunakan fitur revolusioner Gemini AI Auto-Fill, berdiskusi & menerbitkan 'Custom Offer' eksklusif via ChatWidget, menginput resi kurir, menginspeksi foto unit HP lama pada pesanan trade-in, serta merespons ulasan pembeli."),
        ("Pembeli (Buyer / Konsumen)", "Pencari & Pembeli Gadget Terpercaya", 
         "Mengakses katalog dan profil pembeli (/buyer-profile), simulasi harga pasar di AI Valuation, cek IMEI di Scanner, simulasi kebutuhan di Smart Match, mengajukan Trade-In dengan 6 foto kondisi, checkout aman berproteksi Cloudflare Turnstile & Midtrans Escrow, pelacakan resi BinderByte, serta konfirmasi order/retur.")
    ]
    
    actor_col_widths = [Inches(1.8), Inches(1.8), Inches(2.9)]
    for c_idx, h_text in enumerate(actor_headers):
        cell = actor_table.rows[0].cells[c_idx]
        cell.width = actor_col_widths[c_idx]
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        set_cell_background(cell, "1B365D")
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h_text)
        r.font.name = 'Calibri'
        r.font.size = Pt(10)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    for r_idx, (a, b, c) in enumerate(actor_rows):
        row = actor_table.rows[r_idx + 1]
        bg = "FFFFFF" if r_idx % 2 == 0 else "F9FAFC"
        for c_idx, val in enumerate([a, b, c]):
            cell = row.cells[c_idx]
            cell.width = actor_col_widths[c_idx]
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            set_cell_background(cell, bg)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.name = 'Calibri'
            r.font.size = Pt(9.5)
            if c_idx == 0:
                r.bold = True
                r.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
            else:
                r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    p_spacer2 = doc.add_paragraph()
    p_spacer2.paragraph_format.space_before = Pt(6)
    p_spacer2.paragraph_format.space_after = Pt(6)

    # --- BAB 2 ---
    add_custom_heading("2. User Story, User Journey & Skenario Pengujian", level=1)
    doc.add_paragraph(
        "Bagian ini merinci spesifikasi kebutuhan fungsional dalam format User Story standar industri agile, alur perjalanan pengguna "
        "(User Journey), serta kriteria penerimaan yang mencakup Skenario Pengujian Utama (Happy Path / Main Flow) dan Skenario "
        "Edge Cases (Exception / Error Flow) untuk masing-masing aktor pada sistem GadgetTrustX. "
        "Sesuai acuan spesifikasi kebutuhan, seluruh tabel skenario difokuskan secara presisi pada Kondisi Prasyarat (Given) "
        "dan Tindakan Pemicu Pengguna (When)."
    )

    def add_user_story_box(actor_name, story_text, journey_text):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.rows[0].cells[0]
        cell.width = Inches(6.5)
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        set_cell_background(cell, "F0F4F8")
        
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
            f'<w:left w:val="single" w:sz="24" w:space="0" w:color="1B365D"/>'
            f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
            f'<w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        r_head = p.add_run(f"USER STORY & USER JOURNEY – {actor_name.upper()}\n")
        r_head.bold = True
        r_head.font.size = Pt(11)
        r_head.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        
        r_lbl1 = p.add_run("User Story: ")
        r_lbl1.bold = True
        r_lbl1.font.size = Pt(10)
        r_story = p.add_run(f"{story_text}\n\n")
        r_story.font.size = Pt(10)
        
        r_lbl2 = p.add_run("User Journey: ")
        r_lbl2.bold = True
        r_lbl2.font.size = Pt(10)
        r_journey = p.add_run(f"{journey_text}")
        r_journey.font.size = Pt(10)

    def render_scenario_table(headers, data, widths, is_edge_case=False):
        tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl, color="B0C4DE", sz="6")
        
        header_color = "2C5282" if is_edge_case else "1B365D"
        for c_idx, h in enumerate(headers):
            cell = tbl.rows[0].cells[c_idx]
            cell.width = widths[c_idx]
            set_cell_margins(cell, top=90, bottom=90, left=110, right=110)
            set_cell_background(cell, header_color)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(h)
            r.bold = True
            r.font.name = 'Calibri'
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            
        for r_idx, row_vals in enumerate(data):
            row = tbl.rows[r_idx + 1]
            bg = "FFFFFF" if r_idx % 2 == 0 else "F9FAFC"
            for c_idx, val in enumerate(row_vals):
                cell = row.cells[c_idx]
                cell.width = widths[c_idx]
                set_cell_margins(cell, top=90, bottom=90, left=110, right=110)
                set_cell_background(cell, bg)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(val)
                r.font.name = 'Calibri'
                r.font.size = Pt(9)
                if c_idx == 0:
                    r.bold = True
                    r.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
                else:
                    r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
                    
        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(4)
        p_sp.paragraph_format.space_after = Pt(4)

    # --- 2.1 SUPERADMIN ---
    add_custom_heading("2.1 Aktor: Pengelola Platform & Verifikator (Superadmin)", level=2)
    
    admin_story = (
        "Sebagai Superadmin / Inspector GadgetTrustX, saya ingin memantau metrik ekosistem platform, memvalidasi dan memberikan "
        "badge 'Verified by TrustX' pada katalog gadget, menyelesaikan tiket sengketa retur barang (Return Refund Approval), "
        "serta mengelola akun dan reputasi toko penjual agar operasional pasar gadget bekas berlangsung aman, adil, dan kredibel."
    )
    admin_journey = (
        "Superadmin masuk ke Dasbor Admin (/admin) -> Membaca ringkasan total pengguna, produk, transaksi, dan pending return pada tab Dashboard "
        "-> Membuka tab Orders untuk meninjau status pesanan escrow -> Meninjau bukti foto kerusakan/cacat pada tiket komplain retur pembeli "
        "-> Memilih tindakan 'Approve' (dana dikembalikan ke pembeli) atau 'Reject' (dana diteruskan ke penjual) -> Membuka tab Produk "
        "untuk meninjau listing baru dan mengaktifkan tanda verifikasi 'Verified by TrustX' -> Membuka tab User untuk memverifikasi akun toko "
        "atau menyematkan badge reputasi (Official Store / Top Seller)."
    )
    add_user_story_box("Superadmin / TrustX Inspector", admin_story, admin_journey)
    
    add_custom_heading("A. Skenario Pengujian Utama (Main Acceptance Criteria)", level=3)
    admin_main_headers = ["No", "Nama Skenario", "Kondisi Awal (Given)", "Tindakan (When)"]
    admin_main_widths = [Inches(0.5), Inches(1.8), Inches(2.1), Inches(2.1)]
    admin_main_data = [
        ("1", "Pemberian Badge Verified Produk", 
         "Superadmin berada di tab Produk pada dasbor admin dan memilih produk yang telah lolos audit.", 
         "Menekan tombol 'Verifikasi' untuk mengaktifkan status 'Verified by TrustX' pada listing gadget."),
        ("2", "Penyelesaian Sengketa Retur (Approve)", 
         "Terdapat pengajuan retur pesanan berstatus 'Pending' dengan lampiran foto fisik rusak dari pembeli.", 
         "Menekan tombol 'Approve Return' pada modal konfirmasi untuk mengembalikan dana ke saldo pembeli."),
        ("3", "Penolakan Komplain Retur (Reject)", 
         "Bukti foto retur yang diunggah pembeli tidak sesuai dengan laporan kerusakan pada unboxing.", 
         "Menekan tombol 'Reject Return' untuk membatalkan klaim dan meneruskan dana escrow ke penjual."),
        ("4", "Pemberian Badge Reputasi Toko", 
         "Superadmin berada di tab User dan membuka detail penjual dengan riwayat transaksi tinggi.", 
         "Mengaktifkan badge 'Official Store' dan 'Top Seller' pada profil akun penjual tersebut.")
    ]
    render_scenario_table(admin_main_headers, admin_main_data, admin_main_widths, is_edge_case=False)

    add_custom_heading("B. Skenario Edge Cases & Penanganan Pengecualian", level=3)
    admin_edge_headers = ["No", "Skenario Edge Case", "Kondisi Awal (Given)", "Pemicu (When)"]
    admin_edge_widths = [Inches(0.5), Inches(1.8), Inches(2.1), Inches(2.1)]
    admin_edge_data = [
        ("1", "Pemberian Verifikasi pada Produk yang Dihapus", 
         "Penjual menghapus produknya sesaat sebelum admin menekan tombol verifikasi.", 
         "Superadmin menekan tombol verifikasi pada baris produk yang datanya telah terhapus di Supabase."),
        ("2", "Aproval Retur Bersamaan dengan Konfirmasi Selesai", 
         "Pembeli mengajukan komplain retur, namun secara tidak sengaja menekan 'Pesanan Selesai' di perangkat lain.", 
         "Superadmin mengeksekusi 'Approve Return' tepat saat status pesanan berubah menjadi 'Completed'."),
        ("3", "Penghapusan User dengan Transaksi Escrow Berjalan", 
         "User penjual memiliki 3 pesanan berstatus 'Processing' yang dananya masih tersimpan di rekening escrow.", 
         "Superadmin menekan tombol hapus user pada akun penjual tersebut di tab User Management.")
    ]
    render_scenario_table(admin_edge_headers, admin_edge_data, admin_edge_widths, is_edge_case=True)

    # --- 2.2 SELLER ---
    add_custom_heading("2.2 Aktor: Penjual Gadget (Seller / Toko Mitra)", level=2)
    
    seller_story = (
        "Sebagai Penjual Gadget, saya ingin mendaftarkan identitas toko, mengunggah katalog gadget dengan spesifikasi otomatis (Gemini AI Auto-Fill), "
        "menawarkan harga nego khusus melalui ChatWidget (Custom Offer), mencetak label & menginput resi pengiriman, menginspeksi gadget lama pada transaksi tukar tambah, "
        "serta merespons ulasan pembeli agar dagangan saya laku cepat dengan kepastian dana terlindungi escrow."
    )
    seller_journey = (
        "Penjual membuka dasbor toko (/seller-profile) -> Mengisi identitas toko dan titik lokasi gudang pengiriman via AddressPicker "
        "-> Membuka tab Katalog dan menekan tombol 'Tambah Gadget' -> Mengetik nama model dan menekan tombol 'Auto-Fill Spesifikasi AI' "
        "untuk mengisi otomatis chipset, RAM, storage, dan deskripsi produk bertenaga Gemini AI -> Mengunggah foto produk ke Supabase Storage "
        "-> Menerima pesan negosiasi dari calon pembeli via ChatWidget -> Menerbitkan kartu 'Custom Offer' eksklusif dengan harga mufakat "
        "-> Menerima notifikasi pesanan telah dibayar ke Escrow -> Mengemas barang, memilih kurir, dan menginput nomor resi pada modal pengiriman "
        "-> Membuka tab Trade-In untuk melacak paket HP lama pembeli, mengunggah foto inspeksi fisik, dan mengonfirmasi penerimaan "
        "-> Membalas ulasan dan foto ulasan pembeli di tab Ulasan."
    )
    add_user_story_box("Penjual Gadget (Seller)", seller_story, seller_journey)

    add_custom_heading("A. Skenario Pengujian Utama (Main Acceptance Criteria)", level=3)
    seller_main_headers = ["No", "Nama Skenario", "Kondisi Awal (Given)", "Tindakan (When)"]
    seller_main_widths = [Inches(0.5), Inches(1.8), Inches(2.1), Inches(2.1)]
    seller_main_data = [
        ("1", "Otomatisasi Spesifikasi Produk (AI Auto-Fill)", 
         "Penjual membuka formulir tambah produk dan mengetik 'Samsung Galaxy S24 Ultra 256GB'.", 
         "Menekan tombol 'Auto-Fill Spesifikasi AI' untuk mengisi otomatis chipset Snapdragon 8 Gen 3, RAM 12GB, dan deskripsi."),
        ("2", "Pembuatan Penawaran Khusus (Custom Offer)", 
         "Penjual dan pembeli sedang bernegosiasi harga di jendela ChatWidget.", 
         "Penjual mengisi nominal harga Rp 8.200.000 dan menekan tombol 'Kirim Custom Offer' yang terkunci khusus untuk pembeli tersebut."),
        ("3", "Input Nomor Resi & Kurir Pengiriman", 
         "Pesanan pembeli telah dibayar lunas ke rekening escrow dan berstatus 'Processing'.", 
         "Penjual membuka modal 'Kirim Barang', memilih ekspedisi 'JNE Express', dan menginput nomor resi 'JNE882910293'."),
        ("4", "Inspeksi Fisik Unit Tukar Tambah (Trade-In)", 
         "Penjual menerima paket kiriman HP lama dari pembeli pada transaksi tukar tambah.", 
         "Membuka modal inspeksi di tab Trade-In, mengunggah foto fisik unit, dan menekan 'Konfirmasi Selesai & Terima Unit'."),
        ("5", "Merespons Ulasan Pembeli", 
         "Pembeli memberikan ulasan bintang 5 dengan foto barang sampai di tab Ulasan.", 
         "Penjual mengetikkan teks ucapan terima kasih dan menekan tombol 'Kirim Tanggapan Penjual'.")
    ]
    render_scenario_table(seller_main_headers, seller_main_data, seller_main_widths, is_edge_case=False)

    add_custom_heading("B. Skenario Edge Cases & Penanganan Pengecualian", level=3)
    seller_edge_headers = ["No", "Skenario Edge Case", "Kondisi Awal (Given)", "Pemicu (When)"]
    seller_edge_widths = [Inches(0.5), Inches(1.8), Inches(2.1), Inches(2.1)]
    seller_edge_data = [
        ("1", "Input Stok atau Harga Tidak Logis", 
         "Penjual mengisi formulir pendaftaran produk gadget baru.", 
         "Memasukkan nilai harga '0' atau nilai stok '-5' lalu menekan tombol simpan."),
        ("2", "Pengunggahan Foto Non-Gambar pada Produk", 
         "Penjual berada di bagian unggah foto produk gadget.", 
         "Memilih file dokumen PDF atau file executable yang diubah ekstensinya menjadi .png."),
        ("3", "Penolakan Unit Trade-In dengan Bukti Cacat", 
         "HP lama yang dikirim pembeli mengalami layar retak total (berbeda dari klaim formulir 'Mulus').", 
         "Penjual membuka modal 'Tolak Trade-In', memilih kategori cacat layar, mengunggah foto bukti, dan memasukkan resi retur.")
    ]
    render_scenario_table(seller_edge_headers, seller_edge_data, seller_edge_widths, is_edge_case=True)

    # --- 2.3 BUYER ---
    add_custom_heading("2.3 Aktor: Pembeli Gadget (Buyer / Konsumen)", level=2)
    
    buyer_story = (
        "Sebagai Pembeli Gadget, saya ingin menjelajahi katalog terverifikasi, memvalidasi harga via AI Valuation, mengecek legalitas IMEI di Scanner, "
        "memperoleh rekomendasi di Smart Match, mengajukan Tukar Tambah dengan potongan instan, membayar aman via Midtrans Escrow, "
        "melacak posisi resi real-time, serta memberikan ulasan agar mendapatkan gadget impian dengan aman tanpa takut tertipu."
    )
    buyer_journey = (
        "Pembeli membuka website GadgetTrustX -> Menggunakan menu AI Valuation (/price-checker) untuk melihat harga pasar gadget second "
        "-> Membuka Scanner (/verification) untuk mengecek IMEI hp incaran -> Menjalankan wizard Smart Match (/smart-matching) "
        "untuk memperoleh rekomendasi gadget terbaik -> Membuka katalog Marketplace (/marketplace) dan memilih gadget berlabel 'Verified by TrustX' "
        "-> Membuka halaman Trade-In (/trade-in), mengunggah 6 sudut foto hp lama, dan memperoleh nilai potongan harga instan "
        "-> Menuju halaman Checkout (/checkout), memilih kurir JNE/SiCepat/Pos, menyelesaikan verifikasi Cloudflare Turnstile, "
        "dan membayar via QRIS GoPay / Virtual Account Midtrans -> Membuka dasbor /buyer-profile untuk memantau countdown pembayaran dan melacak resi pengiriman real-time "
        "-> Menerima unit gadget dan menekan 'Konfirmasi Selesai' atau mengajukan retur jika cacat -> Memberikan ulasan dan rating bintang 5."
    )
    add_user_story_box("Pembeli Gadget (Buyer)", buyer_story, buyer_journey)

    add_custom_heading("A. Skenario Pengujian Utama (Main Acceptance Criteria)", level=3)
    buyer_main_headers = ["No", "Nama Skenario", "Kondisi Awal (Given)", "Tindakan (When)"]
    buyer_main_widths = [Inches(0.5), Inches(1.8), Inches(2.1), Inches(2.1)]
    buyer_main_data = [
        ("1", "Validasi IMEI via TrustX Scanner", 
         "Pembeli berada di halaman Scanner (/verification) dan memasukkan 15 digit IMEI valid.", 
         "Menekan tombol 'Verifikasi Keamanan Perangkat' untuk mencocokkan database TAC global dan cek status blacklist."),
        ("2", "Rekomendasi Cerdas via Smart Match", 
         "Pembeli menyelesaikan 3 langkah preferensi budget Rp 10.000.000 dan fokus 'Gaming' di /smart-matching.", 
         "Menekan tombol 'Analisis Rekomendasi' untuk memicu Gemini AI menghitung match score dan reasoning spesifikasi."),
        ("3", "Pengajuan Tukar Tambah 6 Sudut Foto", 
         "Pembeli berada di halaman /trade-in pada tab formulir tukar tambah HP lama.", 
         "Mengisi spesifikasi, mengunggah foto 6 sudut (Depan, Belakang, Layar, Samping, Box, Cacat), dan menekan 'Hitung Valuasi AI'."),
        ("4", "Checkout Aman dengan Midtrans Escrow", 
         "Pembeli berada di halaman /checkout dengan barang pilihan dan diskon trade-in aktif.", 
         "Menyelesaikan CAPTCHA Cloudflare Turnstile dan menekan tombol 'Bayar Sekarang' untuk membuka modal pembayaran Midtrans."),
        ("5", "Pelacakan Pengiriman Resi Real-Time", 
         "Penjual telah menginput nomor resi pengiriman pada pesanan pembeli.", 
         "Pembeli menekan tombol 'Lacak Pengiriman' di dasbor /buyer-profile untuk melihat timeline ekspedisi via API BinderByte."),
        ("6", "Pengajuan Retur Barang Bermasalah", 
         "Paket diterima pembeli namun kondisi fisik tidak sesuai deskripsi (misal tombol power rusak).", 
         "Membuka modal 'Ajukan Retur' di /buyer-profile, mengisi alasan kerusakan, mengunggah foto bukti, dan menekan submit.")
    ]
    render_scenario_table(buyer_main_headers, buyer_main_data, buyer_main_widths, is_edge_case=False)

    add_custom_heading("B. Skenario Edge Cases & Penanganan Kondisi Invalid", level=3)
    buyer_edge_headers = ["No", "Skenario Edge Case", "Kondisi Awal (Given)", "Pemicu (When)"]
    buyer_edge_widths = [Inches(0.5), Inches(1.8), Inches(2.1), Inches(2.1)]
    buyer_edge_data = [
        ("1", "Input Nomor IMEI Kurang dari 15 Digit", 
         "Pembeli berada di kolom input halaman Scanner (/verification).", 
         "Memasukkan nomor IMEI yang hanya berisi 12 digit lalu menekan enter."),
        ("2", "Kedaluwarsa Batas Waktu Pembayaran (Expired)", 
         "Sesi pembayaran GoPay QRIS (15 menit) atau Virtual Account (24 jam) telah habis pada modal pembayaran.", 
         "Pembeli melakukan transfer setelah live countdown timer menunjukkan angka 00:00."),
        ("3", "Blokir Akses Bot Cloudflare Turnstile", 
         "Percobaan checkout dilakukan tanpa menyelesaikan token tantangan keamanan Turnstile.", 
         "Pengguna atau script otomatis menekan tombol 'Bayar Sekarang' dengan token CAPTCHA kosong atau invalid."),
        ("4", "Pembelian Melebihi Batas Stok yang Tersedia", 
         "Stok unit gadget hanya tersisa 1 buah di halaman detail produk.", 
         "Pembeli mencoba menambahkan kuantitas sebanyak 2 buah ke dalam keranjang belanja."),
        ("5", "Pengajuan Komplain Melewati Garansi 72 Jam", 
         "Pesanan telah berstatus 'Delivered' selama lebih dari 3 hari (72 jam).", 
         "Pembeli mencoba membuka opsi 'Ajukan Retur' yang tombolnya telah otomatis terkunci oleh sistem.")
    ]
    render_scenario_table(buyer_edge_headers, buyer_edge_data, buyer_edge_widths, is_edge_case=True)

    # --- BAB 3 ---
    add_custom_heading("3. Aturan Bisnis Khusus (High-Level Business Rules)", level=1)
    doc.add_paragraph(
        "Guna menjamin integritas ekosistem marketplace, kepastian hak pembeli, kelancaran perputaran modal penjual, "
        "serta kepatuhan regulasi telekomunikasi dan perlindungan konsumen, platform GadgetTrustX menerapkan aturan bisnis "
        "tingkat tinggi (High-Level Business Rules) berikut:"
    )

    rules = [
        ("3.1 Perlindungan Dana Escrow & Auto-Complete 3x24 Jam", 
         "Seluruh dana pembayaran pembeli ditahan secara aman di rekening penampungan sementara (TrustX Escrow Holding). "
         "Dana dilarang dicairkan ke saldo penjual sebelum pembeli menekan 'Konfirmasi Selesai' atau melewati batas waktu "
         "perlindungan otomatis 3 x 24 jam (72 jam) sejak nomor resi kurir terkonfirmasi 'Delivered'. Aturan ini mengunci hak komplain "
         "pembeli sekaligus menjamin kepastian pencairan hak penjual tanpa penahanan dana sepihak."),
        
        ("3.2 Dual-Check Trade-In Lock & Cut-off Diskon Valuasi", 
         "Potongan harga dari taksiran Google Gemini AI pada transaksi tukar tambah bersifat diskon bersyarat yang terikat pada ID trade-in. "
         "Unit gadget baru yang dibeli baru dapat dikirimkan oleh penjual setelah paket unit lama pembeli tiba di gudang penjual "
         "dan lolos inspeksi fisik 6 sudut foto. Jika fisik unit lama cacat parah atau tidak sesuai formulir, penjual berhak "
         "mengajukan tawar balik (Counter Offer) atau menolak trade-in dan membatalkan pesanan."),
        
        ("3.3 Validasi IMEI Unik & Anti-Peredaran Barang Ilegal", 
         "Setiap unit gadget yang dipublikasikan di katalog marketplace wajib mencantumkan 15 digit nomor IMEI yang lolos uji checksum Luhn. "
         "Sistem secara otomatis menolak pendaftaran IMEI yang berstatus duplikat di database aktif guna mencegah peredaran unit kloningan, "
         "ponsel curian, maupun listing ganda oleh pihak tidak bertanggung jawab."),
        
        ("3.4 Proteksi Bot, Idempotency Token & Rate-Limiting Turnstile", 
         "Setiap transaksi checkout dan autentikasi diproteksi oleh token Cloudflare Turnstile non-intrusif pada lapisan frontend "
         "dan verifikasi serverless di sisi backend. Setiap permintaan mutasi transaksi menyertakan idempotency key untuk mencegah "
         "terjadinya pemotongan saldo ganda atau eksploitasi race condition saat jaringan tidak stabil.")
    ]

    for title, desc in rules:
        p_r = doc.add_paragraph()
        p_r.paragraph_format.space_before = Pt(4)
        p_r.paragraph_format.space_after = Pt(4)
        r_t = p_r.add_run(f"{title}: ")
        r_t.bold = True
        r_t.font.name = 'Calibri'
        r_t.font.size = Pt(10.5)
        r_t.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        r_d = p_r.add_run(desc)
        r_d.font.name = 'Calibri'
        r_d.font.size = Pt(10.5)

    # --- BAB 4 ---
    add_custom_heading("4. Kebutuhan Non-Fungsional & Standar Pengalaman Pengguna (UX)", level=1)
    doc.add_paragraph(
        "Selain pemenuhan kebutuhan fungsional, platform GadgetTrustX dibangun dengan mengacu pada standar rekayasa perangkat lunak modern, "
        "keamanan data berlapis, dan kenyamanan antarmuka pengguna (User Experience):"
    )

    nfrs = [
        ("Keandalan Akses & Arsitektur Serverless Next.js Turbopack", 
         "Platform dibangun di atas Next.js 15 App Router dan Turbopack yang dideploy ke jaringan Netlify Edge CDN. Waktu respon pemuatan halaman awal (First Contentful Paint) berada di bawah 1.5 detik dengan ketersediaan layanan (Uptime SLA) mencapai 99.9%."),
        ("Keamanan Multi-Tenant & Row Level Security (RLS) PostgreSQL Supabase", 
         "Sebanyak 10 tabel inti (profiles, devices, reviews, cart_items, orders, order_items, wtb_listings, chats, chat_messages, trade_in_requests) diproteksi secara ketat menggunakan Row Level Security (RLS) di lapisan database engine. Pengguna hanya memiliki hak akses terhadap entitas data miliknya sendiri."),
        ("Ketahanan & Fail-Safe Multi-Key Rotation Pool Gemini AI", 
         "Modul integrasi AI mengimplementasikan key rotation pool dengan 7 API keys yang diputar secara acak disertai mekanisme caching memori. Hal ini menjamin fitur AI Auto-Fill, AI Valuation, dan Smart Match tetap beroperasi normal tanpa kendala kuota harian."),
        ("Penyimpanan Media Terpusat & Optimalisasi Gambar", 
         "Seluruh foto produk, foto inspeksi trade-in, foto ulasan, dan avatar disimpan pada Supabase Storage Bucket 'device-media' yang terintegrasi dengan kebijakan akses aman dan kompresi visual otomatis."),
        ("Desain Responsif & Transisi Glassmorphism Halus", 
         "Antarmuka menerapkan paradigma Mobile-First dengan tema Dark Futuristic Glassmorphism berbasis Tailwind CSS dan Framer Motion, memastikan tampilan tetap proporsional dan nyaman diakses dari smartphone, tablet, maupun layar desktop."),
        ("Pesan Kesalahan Ramah Pengguna & Fail-Safe Fallback", 
         "Sistem dilarang memunculkan kode galat mentah (seperti pesan SQL error atau raw code 500). Setiap kegagalan koneksi atau pembayaran ditransformasikan menjadi dialog instruksi yang solutif dalam Bahasa Indonesia.")
    ]

    for title, desc in nfrs:
        p_n = doc.add_paragraph(style='List Bullet')
        p_n.paragraph_format.space_before = Pt(3)
        p_n.paragraph_format.space_after = Pt(3)
        r_nt = p_n.add_run(f"{title}: ")
        r_nt.bold = True
        r_nt.font.name = 'Calibri'
        r_nt.font.size = Pt(10)
        r_nt.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        r_nd = p_n.add_run(desc)
        r_nd.font.name = 'Calibri'
        r_nd.font.size = Pt(10)

    # Save to both paths
    out_dir_downloads = r'C:\Users\LEANDRA\Downloads'
    out_file_downloads = os.path.join(out_dir_downloads, 'Spesifikasi Kebutuhan Perangkat Lunak - GadgetTrustX.docx')
    doc.save(out_file_downloads)
    
    out_file_workspace = r'f:\smart-device-marketplace\Spesifikasi Kebutuhan Perangkat Lunak - GadgetTrustX.docx'
    doc.save(out_file_workspace)
    
    print(f'Document successfully created:\n1. {out_file_downloads}\n2. {out_file_workspace}')

if __name__ == '__main__':
    create_document()
